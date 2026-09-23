"""Deliverable document generation tools for Sovereign AI Workbench.

Implements 02_DESIGN_DOC.md §10 and 03_SECURITY_AND_ACCESS.md §11:
- create_docx: template substitution ({{field}}), table formatting, placeholder verification.
- create_xlsx: multi-sheet workbook generation with SEC-15 formula sanitization.
- create_pptx: corporate presentation generation with structured slide content.
- create_calculation_report: formal engineering calculation report in DOCX and XLSX formats.
- Database artifact registration and SHA-256 integrity hash tracking.
"""

from datetime import datetime, timezone
import hashlib
import logging
from pathlib import Path
import re
from typing import Any, Dict, List, Literal, Optional, Tuple, Union

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from pydantic import BaseModel, Field
import pptx

from backend.core.config import get_settings
from backend.core.db import Artifact, Task, User, get_session_factory
from backend.core.paths import safe_path
from tools.registry import ToolContext, audited_tool

logger = logging.getLogger("sovereign-workbench.tools.documents")

# Dangerous Excel formula keywords per 03_SECURITY_AND_ACCESS.md §11
DANGEROUS_FORMULA_PATTERNS = [
    r"\bWEBSERVICE\b",
    r"\bHYPERLINK\s*\(\s*[\"']https?://",
    r"\bEXEC\b",
    r"\bCMD\b",
    r"\[.*\.xlsx?\]",  # External workbook references
    r"\bSHELL\b",
]

ALLOWED_FORMULA_PREFIXES = (
    "=SUM(",
    "=AVERAGE(",
    "=COUNT(",
    "=MIN(",
    "=MAX(",
    "=ROUND(",
    "=IF(",
    "=(",
)


def compute_file_sha256(file_path: Path) -> str:
    """Compute SHA-256 hash of a file on disk."""
    hasher = hashlib.sha256()
    with open(file_path, "rb") as f:
        while chunk := f.read(65536):
            hasher.update(chunk)
    return hasher.hexdigest()


def register_artifact_in_db(
    run_id: str,
    user_id: str,
    kind: str,
    filename: str,
    stored_path: str,
    sha256_hash: str,
) -> str:
    """Record generated deliverable in SQLite artifacts table with integrity checksum."""
    artifact_id = f"art-{run_id}-{kind}-{int(datetime.now(timezone.utc).timestamp())}"
    factory = get_session_factory()

    try:
        with factory() as db:
            # Ensure task exists to satisfy foreign key if present
            task = db.query(Task).filter(Task.id == run_id).first()
            if not task:
                # Ensure user exists
                user = db.query(User).filter(User.id == user_id).first()
                if not user:
                    user = User(
                        id=user_id,
                        username=f"user_{user_id}",
                        password_hash="mock_hash",
                        role="engineer",
                        clearance=1,
                    )
                    db.add(user)
                    db.flush()

                task = Task(
                    id=run_id,
                    owner_id=user_id,
                    request_text="Deliverable generation task",
                    status="completed",
                )
                db.add(task)
                db.flush()

            artifact = Artifact(
                id=artifact_id,
                task_id=run_id,
                owner_id=user_id,
                kind=kind,
                filename=filename,
                stored_path=stored_path,
                sha256=sha256_hash,
                status="PENDING_REVIEW",
            )
            db.add(artifact)
            db.commit()
    except Exception as e:
        logger.warning(f"Could not register artifact in DB: {e}")

    return artifact_id


# ==============================================================================
# 1. DOCX Generator
# ==============================================================================

class CreateDocxArgs(BaseModel):
    """Input parameters for create_docx tool."""

    template_path: str = Field(
        default="templates/approval_note.docx",
        description="Path to Word template containing {{placeholders}}",
    )
    fields: Dict[str, Any] = Field(
        default_factory=dict,
        description="Dictionary of placeholder keys and replacement values",
    )
    filename: Optional[str] = Field(default=None, description="Output filename (defaults to template name)")


class CreateDocxResult(BaseModel):
    """Result of create_docx tool execution."""

    artifact_id: str
    path: str
    filename: str
    sha256: str
    fields_replaced: List[str]
    unresolved_placeholders: List[str]


def _replace_text_in_paragraph(paragraph: docx.text.paragraph.Paragraph, fields: Dict[str, Any]) -> List[str]:
    """Replace {{key}} tokens across text runs in a paragraph."""
    replaced = []
    full_text = paragraph.text
    if "{{" not in full_text:
        return replaced

    for key, val in fields.items():
        pattern = f"{{{{{key}}}}}"
        if pattern in full_text:
            str_val = str(val) if not isinstance(val, (list, dict)) else "\n".join(str(item) for item in val)
            full_text = full_text.replace(pattern, str_val)
            replaced.append(key)

    if replaced:
        paragraph.text = full_text

    return replaced


@audited_tool(name="create_docx", side_effects=True, needs_role=None)
def create_docx(args: CreateDocxArgs, ctx: ToolContext) -> CreateDocxResult:
    """Generate Word document deliverable from template with placeholder substitution."""
    settings = get_settings()
    template_p = Path(args.template_path)
    if not template_p.exists():
        raise FileNotFoundError(f"DOCX template not found at '{template_p}'.")

    # Output directory
    output_dir = Path(settings.ALLOWED_OUTPUT_DIRS) / ctx.run_id
    output_dir.mkdir(parents=True, exist_ok=True)

    filename = args.filename or template_p.name
    if not filename.lower().endswith(".docx"):
        filename += ".docx"
    output_path = output_dir / filename

    doc = docx.Document(str(template_p))
    replaced_keys = []

    # Merge context defaults
    effective_fields = dict(args.fields)
    effective_fields.setdefault("run_id", ctx.run_id)
    effective_fields.setdefault("prepared_by", ctx.user_id)
    effective_fields.setdefault("clearance", ctx.clearance)
    effective_fields.setdefault("date", datetime.now(timezone.utc).strftime("%Y-%m-%d"))

    # Replace in body paragraphs
    for p in doc.paragraphs:
        r = _replace_text_in_paragraph(p, effective_fields)
        replaced_keys.extend(r)

    # Replace in tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    r = _replace_text_in_paragraph(p, effective_fields)
                    replaced_keys.extend(r)

    # Detect any remaining unresolved placeholders
    remaining = []
    all_text = "\n".join(p.text for p in doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                all_text += "\n" + "\n".join(p.text for p in cell.paragraphs)

    unresolved_matches = re.findall(r"\{\{([a-zA-Z0-9_\-]+)\}\}", all_text)
    remaining = list(set(unresolved_matches))

    doc.save(str(output_path))
    sha256 = compute_file_sha256(output_path)
    art_id = register_artifact_in_db(
        run_id=ctx.run_id,
        user_id=ctx.user_id,
        kind="docx",
        filename=filename,
        stored_path=str(output_path),
        sha256_hash=sha256,
    )

    return CreateDocxResult(
        artifact_id=art_id,
        path=str(output_path),
        filename=filename,
        sha256=sha256,
        fields_replaced=sorted(list(set(replaced_keys))),
        unresolved_placeholders=remaining,
    )


# ==============================================================================
# 2. XLSX Generator with Formula Sanitization (SEC-15)
# ==============================================================================

class CreateXlsxArgs(BaseModel):
    """Input parameters for create_xlsx tool."""

    sheets: Dict[str, List[List[Any]]] = Field(
        default_factory=dict,
        description="Dictionary mapping sheet names (e.g. 'Data', 'Calculations') to 2D row/column data.",
    )
    filename: Optional[str] = Field(default="analysis.xlsx", description="Output filename")


class CreateXlsxResult(BaseModel):
    """Result of create_xlsx tool execution."""

    artifact_id: str
    path: str
    filename: str
    sha256: str
    sheets_created: List[str]
    sanitized_cells_count: int


def sanitize_spreadsheet_cell(val: Any) -> Tuple[Any, bool]:
    """Sanitize spreadsheet cell value per SEC-15 formula injection defenses.

    Values starting with = + - @ from untrusted sources are escaped as text unless
    explicitly matched to an approved safe formula. Blocks DDE, WEBSERVICE, and external links.
    """
    if not isinstance(val, str):
        return val, False

    s = val.strip()
    if not s:
        return val, False

    # Check for dangerous formulas (DDE, WEBSERVICE, external links)
    for pattern in DANGEROUS_FORMULA_PATTERNS:
        if re.search(pattern, s, re.IGNORECASE):
            # Escape as safe string by prepending single quote
            return f"'{s}", True

    # If it starts with '=', verify whether it is an allowed safe formula
    if s.startswith("="):
        is_safe_formula = any(s.upper().startswith(prefix) for prefix in ALLOWED_FORMULA_PREFIXES) or bool(
            re.match(r"^=[A-Z0-9_]+(\s*[\+\-\*\/\^]\s*[A-Z0-9_]+)*$", s, re.IGNORECASE)
        )
        if is_safe_formula:
            return s, False
        # Untrusted formula -> escape as text
        return f"'{s}", True

    # If it starts with + - @ and is not purely a signed number
    if s[0] in ("+", "-", "@"):
        try:
            float(s)
            return float(s) if "." in s else int(s), False
        except ValueError:
            # Text starting with trigger characters -> escape as text
            return f"'{s}", True

    return val, False


@audited_tool(name="create_xlsx", side_effects=True, needs_role=None)
def create_xlsx(args: CreateXlsxArgs, ctx: ToolContext) -> CreateXlsxResult:
    """Generate Excel workbook deliverable with multi-sheet structure and SEC-15 formula protection."""
    settings = get_settings()
    output_dir = Path(settings.ALLOWED_OUTPUT_DIRS) / ctx.run_id
    output_dir.mkdir(parents=True, exist_ok=True)

    filename = args.filename or "analysis.xlsx"
    if not filename.lower().endswith(".xlsx"):
        filename += ".xlsx"
    output_path = output_dir / filename

    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    # Standard sheet structure
    sheets_data = dict(args.sheets)
    if not sheets_data:
        sheets_data = {
            "Data": [["Item", "Value"]],
            "Calculations": [["Formula", "Result"]],
            "Summary": [["Metric", "Assessment"]],
            "Sources": [["Document", "Page"]],
        }

    header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=10)
    thin_border = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0"),
    )

    sanitized_count = 0
    created_sheets = []

    for sheet_name, rows in sheets_data.items():
        ws = wb.create_sheet(title=sheet_name[:31])  # Excel 31 char limit
        created_sheets.append(ws.title)

        for row_idx, row_vals in enumerate(rows, start=1):
            for col_idx, raw_val in enumerate(row_vals, start=1):
                safe_val, was_sanitized = sanitize_spreadsheet_cell(raw_val)
                if was_sanitized:
                    sanitized_count += 1

                cell = ws.cell(row=row_idx, column=col_idx, value=safe_val)
                cell.border = thin_border

                if row_idx == 1:
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.font = data_font
                    if isinstance(safe_val, (int, float)):
                        cell.alignment = Alignment(horizontal="right")
                    else:
                        cell.alignment = Alignment(horizontal="left")

        # Auto-adjust column widths
        for col in ws.columns:
            max_len = max(len(str(c.value or "")) for c in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    wb.save(str(output_path))
    sha256 = compute_file_sha256(output_path)
    art_id = register_artifact_in_db(
        run_id=ctx.run_id,
        user_id=ctx.user_id,
        kind="xlsx",
        filename=filename,
        stored_path=str(output_path),
        sha256_hash=sha256,
    )

    return CreateXlsxResult(
        artifact_id=art_id,
        path=str(output_path),
        filename=filename,
        sha256=sha256,
        sheets_created=created_sheets,
        sanitized_cells_count=sanitized_count,
    )


# ==============================================================================
# 3. PPTX Generator
# ==============================================================================

class CreatePptxArgs(BaseModel):
    """Input parameters for create_pptx tool."""

    template_path: str = Field(default="templates/presentation.pptx", description="Path to presentation template")
    fields: Dict[str, str] = Field(
        default_factory=dict,
        description="Placeholder tokens mapped to slide values (title, subtitle, agenda, content_bullets, sources)",
    )
    filename: Optional[str] = Field(default="presentation.pptx", description="Output presentation filename")


class CreatePptxResult(BaseModel):
    """Result of create_pptx tool execution."""

    artifact_id: str
    path: str
    filename: str
    sha256: str
    slides_count: int


@audited_tool(name="create_pptx", side_effects=True, needs_role=None)
def create_pptx(args: CreatePptxArgs, ctx: ToolContext) -> CreatePptxResult:
    """Generate presentation slide deck from template with text substitution."""
    settings = get_settings()
    template_p = Path(args.template_path)
    if not template_p.exists():
        raise FileNotFoundError(f"PPTX template not found at '{template_p}'.")

    output_dir = Path(settings.ALLOWED_OUTPUT_DIRS) / ctx.run_id
    output_dir.mkdir(parents=True, exist_ok=True)

    filename = args.filename or "presentation.pptx"
    if not filename.lower().endswith(".pptx"):
        filename += ".pptx"
    output_path = output_dir / filename

    prs = pptx.Presentation(str(template_p))

    effective_fields = dict(args.fields)
    effective_fields.setdefault("author", ctx.user_id)
    effective_fields.setdefault("date", datetime.now(timezone.utc).strftime("%Y-%m-%d"))

    # Replace placeholders across all slide text frames
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                tf = shape.text_frame
                for p in tf.paragraphs:
                    for key, val in effective_fields.items():
                        pattern = f"{{{{{key}}}}}"
                        if pattern in p.text:
                            p.text = p.text.replace(pattern, str(val))

    prs.save(str(output_path))
    sha256 = compute_file_sha256(output_path)
    art_id = register_artifact_in_db(
        run_id=ctx.run_id,
        user_id=ctx.user_id,
        kind="pptx",
        filename=filename,
        stored_path=str(output_path),
        sha256_hash=sha256,
    )

    return CreatePptxResult(
        artifact_id=art_id,
        path=str(output_path),
        filename=filename,
        sha256=sha256,
        slides_count=len(prs.slides),
    )


# ==============================================================================
# 4. Calculation Report Generator (DOCX & XLSX)
# ==============================================================================

class CalculationEntry(BaseModel):
    name: str
    formula: str
    inputs: Dict[str, Union[int, float]]
    result: Union[int, float]
    units: Optional[str] = None
    assumptions: Optional[str] = None


class CreateCalculationReportArgs(BaseModel):
    title: str = "Engineering Calculation Verification Report"
    calculations: List[CalculationEntry]
    filename_base: str = "calculation_report"


class CreateCalculationReportResult(BaseModel):
    docx_path: str
    xlsx_path: str
    calculations_count: int


@audited_tool(name="create_calculation_report", side_effects=True, needs_role=None)
def create_calculation_report(args: CreateCalculationReportArgs, ctx: ToolContext) -> CreateCalculationReportResult:
    """Generate formal engineering calculation report in both DOCX and XLSX formats."""
    settings = get_settings()
    output_dir = Path(settings.ALLOWED_OUTPUT_DIRS) / ctx.run_id
    output_dir.mkdir(parents=True, exist_ok=True)

    # 1. Generate DOCX Report
    docx_path = output_dir / f"{args.filename_base}.docx"
    doc = docx.Document()
    doc.add_heading(args.title, level=0)
    meta = doc.add_paragraph()
    meta.add_run("Run ID: ").bold = True
    meta.add_run(f"{ctx.run_id} | ")
    meta.add_run("Author: ").bold = True
    meta.add_run(f"{ctx.user_id} | ")
    meta.add_run("Date: ").bold = True
    meta.add_run(datetime.now(timezone.utc).strftime("%Y-%m-%d"))

    table = doc.add_table(rows=1, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    hdr[0].text = "Item"
    hdr[1].text = "Inputs"
    hdr[2].text = "Formula"
    hdr[3].text = "Result"
    hdr[4].text = "Assumptions"

    xlsx_rows = [["Calculation Item", "Inputs", "Formula", "Result", "Units", "Assumptions"]]

    for calc in args.calculations:
        row = table.add_row().cells
        row[0].text = calc.name
        row[1].text = ", ".join(f"{k}={v}" for k, v in calc.inputs.items())
        row[2].text = calc.formula
        row[3].text = f"{calc.result} {calc.units or ''}".strip()
        row[4].text = calc.assumptions or "Standard conditions"

        xlsx_rows.append([
            calc.name,
            str(calc.inputs),
            calc.formula,
            calc.result,
            calc.units or "",
            calc.assumptions or "",
        ])

    doc.save(str(docx_path))
    register_artifact_in_db(
        run_id=ctx.run_id,
        user_id=ctx.user_id,
        kind="docx",
        filename=docx_path.name,
        stored_path=str(docx_path),
        sha256_hash=compute_file_sha256(docx_path),
    )

    # 2. Generate XLSX Workbook
    xlsx_path = output_dir / f"{args.filename_base}.xlsx"
    create_xlsx(
        CreateXlsxArgs(
            sheets={"Calculations": xlsx_rows, "Summary": [["Total Items", len(args.calculations)]]},
            filename=xlsx_path.name,
        ),
        ctx=ctx,
    )

    return CreateCalculationReportResult(
        docx_path=str(docx_path),
        xlsx_path=str(xlsx_path),
        calculations_count=len(args.calculations),
    )
