"""Unit tests for deliverable validation library (agent/nodes/validate.py).

Verifies integrity re-open checks, template placeholder leftovers,
macro-enabled rejection, SEC-23 external relationship scanning,
and deterministic calculation verification.
"""

from pathlib import Path
import zipfile
import pytest
import docx
import openpyxl
import pptx

from agent.nodes.validate import (
    scan_office_xml_relationships,
    validate_artifact_file,
    validate_calculation_consistency,
    validate_docx_deliverable,
    validate_pptx_deliverable,
    validate_xlsx_deliverable,
)


def test_reopen_integrity_docx(tmp_path):
    doc_path = tmp_path / "valid.docx"
    doc = docx.Document()
    doc.add_heading("Integrity Test", level=1)
    doc.add_paragraph("Clean document with no placeholders.")
    doc.save(doc_path)

    results = validate_docx_deliverable(doc_path)
    reopen_res = next(r for r in results if r.rule == "docx_reopen_integrity")
    assert reopen_res.passed is True

    placeholder_res = next(r for r in results if r.rule == "no_unresolved_placeholders")
    assert placeholder_res.passed is True


def test_leftover_placeholders_detected_docx(tmp_path):
    doc_path = tmp_path / "unresolved.docx"
    doc = docx.Document()
    doc.add_heading("Approval Note Draft", level=1)
    doc.add_paragraph("Equipment ID: {{asset_id}}")
    doc.add_paragraph("Unresolved field: {{missing_placeholder}}")
    doc.save(doc_path)

    results = validate_docx_deliverable(doc_path)
    placeholder_res = next(r for r in results if r.rule == "no_unresolved_placeholders")
    assert placeholder_res.passed is False
    assert "{{asset_id}}" in placeholder_res.detail or "{{missing_placeholder}}" in placeholder_res.detail


def test_leftover_placeholders_detected_pptx(tmp_path):
    ppt_path = tmp_path / "unresolved.pptx"
    prs = pptx.Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Presentation {{unresolved_title}}"
    prs.save(ppt_path)

    results = validate_pptx_deliverable(ppt_path)
    placeholder_res = next(r for r in results if r.rule == "no_unresolved_placeholders")
    assert placeholder_res.passed is False
    assert "{{unresolved_title}}" in placeholder_res.detail


def test_macro_enabled_format_rejected(tmp_path):
    # Test forbidden macro-enabled extensions
    macro_doc = tmp_path / "malicious.docm"
    macro_doc.write_text("dummy binary content")

    macro_xls = tmp_path / "malicious.xlsm"
    macro_xls.write_text("dummy binary content")

    res_doc = validate_artifact_file(macro_doc)
    macro_res_1 = next(r for r in res_doc if r.rule == "no_macro_enabled_formats")
    assert macro_res_1.passed is False

    res_xls = validate_artifact_file(macro_xls)
    macro_res_2 = next(r for r in res_xls if r.rule == "no_macro_enabled_formats")
    assert macro_res_2.passed is False


def test_sec_23_external_relationships_detected(tmp_path):
    # Create a synthetic zip mimicking an Office doc with TargetMode="External" in .rels
    bad_docx = tmp_path / "external_rel.docx"
    
    # Minimal .rels content with an external target
    rels_content = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
    <Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
        <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/attachedTemplate"
            Target="http://evil.com/template.dotm" TargetMode="External"/>
    </Relationships>"""

    with zipfile.ZipFile(bad_docx, "w") as zf:
        zf.writestr("_rels/.rels", rels_content)
        zf.writestr("word/document.xml", "<document/>")

    results = scan_office_xml_relationships(bad_docx)
    ext_res = next(r for r in results if r.rule == "no_external_relationships")
    assert ext_res.passed is False
    assert "evil.com" in ext_res.detail


def test_embedded_ole_object_detected(tmp_path):
    # Create a synthetic zip with an embedded oleObject binary part
    bad_docx = tmp_path / "ole_embedded.docx"
    
    with zipfile.ZipFile(bad_docx, "w") as zf:
        zf.writestr("_rels/.rels", "<Relationships/>")
        zf.writestr("word/embeddings/oleObject1.bin", b"fake binary payload")

    results = scan_office_xml_relationships(bad_docx)
    ole_res = next(r for r in results if r.rule == "no_embedded_ole_objects")
    assert ole_res.passed is False
    assert "oleObject1.bin" in ole_res.detail


def test_approval_note_source_reference_validation(tmp_path):
    doc_path = tmp_path / "approval_note_test.docx"
    doc = docx.Document()
    doc.add_heading("Approval Note", level=1)
    doc.add_paragraph("Human Review: Decision [ ] Date: [ ]")
    doc.save(doc_path)

    # Missing source reference should fail
    bad_meta = {
        "findings": [
            {"finding": "Pitting found", "severity": "High"}  # No source
        ],
        "visual_observations": [
            {"component": "Pipe", "type": "observed", "limitation": "Visual only"}
        ]
    }
    results = validate_docx_deliverable(doc_path, metadata=bad_meta)
    finding_res = next(r for r in results if r.rule == "findings_source_references")
    assert finding_res.passed is False

    # Valid source reference should pass
    good_meta = {
        "findings": [
            {"finding": "Pitting found", "source": "page_3_scan_1", "severity": "High"}
        ],
        "visual_observations": [
            {"component": "Pipe", "type": "observed", "limitation": "Visual only"}
        ]
    }
    results_good = validate_docx_deliverable(doc_path, metadata=good_meta)
    finding_good = next(r for r in results_good if r.rule == "findings_source_references")
    assert finding_good.passed is True


def test_visual_observation_limitation_validation(tmp_path):
    doc_path = tmp_path / "approval_note_vis.docx"
    doc = docx.Document()
    doc.add_heading("Approval Note", level=1)
    doc.add_paragraph("Reviewer: [ ]")
    doc.save(doc_path)

    # Visual observation missing limitation note must fail
    bad_meta = {
        "visual_observations": [
            {"component": "Pipe Joint", "visible_condition": "Crack", "type": "observed"}  # No limitation
        ]
    }
    results = validate_docx_deliverable(doc_path, metadata=bad_meta)
    vis_res = next(r for r in results if r.rule == "visual_observations_limitation")
    assert vis_res.passed is False

    # Valid visual observation with limitation passes
    good_meta = {
        "visual_observations": [
            {
                "component": "Pipe Joint",
                "visible_condition": "Crack",
                "type": "observed",
                "limitation": "Visual observation only; not a certified dimensional measurement",
            }
        ]
    }
    results_good = validate_docx_deliverable(doc_path, metadata=good_meta)
    vis_good = next(r for r in results_good if r.rule == "visual_observations_limitation")
    assert vis_good.passed is True


def test_validate_calculation_consistency():
    formula = "(measured_thickness - min_thickness) / corrosion_rate"
    variables = {
        "measured_thickness": 8.5,
        "min_thickness": 6.0,
        "corrosion_rate": 0.25,
    }

    # Correct result ( (8.5 - 6.0) / 0.25 = 10.0 )
    res_correct = validate_calculation_consistency(formula, variables, 10.0)
    assert res_correct.passed is True

    # Manipulated/incorrect result
    res_wrong = validate_calculation_consistency(formula, variables, 25.0)
    assert res_wrong.passed is False
    assert "Calculation mismatch" in res_wrong.detail
