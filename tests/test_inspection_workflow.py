"""End-to-end integration tests for Demo A inspection workflow (agent/inspection.py).

Verifies:
- Scanned PDF processing via OCR and image extraction.
- Multimodal visual observation extraction with mandatory limitation notes.
- SOP-301 knowledge retrieval via local RAG with clearance filter.
- Deterministic calculation via calculate tool (no LLM arithmetic).
- Inspection_Approval_Note.docx generation from templates/approval_note.docx.
- Factual findings cite source references ([file, page]).
- Recommendations separated from facts and tied to SOP clauses.
- Formal Human Review Gate present and blank.
- Deliverable validation passes.
- Termination at PENDING_REVIEW; agent cannot set APPROVED.
"""

from pathlib import Path
import pytest
import docx

from agent.inspection import run_inspection_workflow
from backend.core.config import get_settings
from backend.core.db import Artifact, Task, User, get_session_factory
from backend.core.security import hash_password
from scripts.make_demo_data import (
    generate_and_ingest_demo_data,
    generate_scanned_inspection_pdf,
)


@pytest.fixture(scope="module")
def setup_demo_a_environment():
    """Ensure synthetic SOPs and scanned inspection PDF exist and are indexed."""
    settings = get_settings()
    incoming_dir = Path("data/incoming")
    incoming_dir.mkdir(parents=True, exist_ok=True)

    pdf_path = incoming_dir / "demo_scanned_inspection_report.pdf"
    if not pdf_path.exists():
        pdf_path, _ = generate_scanned_inspection_pdf(incoming_dir)

    # Ingest synthetic SOPs including SOP-301 (CONFIDENTIAL / clearance 2)
    generate_and_ingest_demo_data(ingest_to_kb=True)

    import uuid
    uid = uuid.uuid4().hex[:8]
    factory = get_session_factory()
    with factory() as db:
        user = User(
            id=f"eng_demo_{uid}",
            username=f"inspector_{uid}",
            password_hash=hash_password("InspectPass123!"),
            role="engineer",
            clearance=2,  # CONFIDENTIAL
            active=True,
        )
        db.add(user)
        db.commit()

        task = Task(
            id=f"task_demo_{uid}",
            owner_id=user.id,
            request_text="Process scanned inspection report for PV-402 and generate approval note deliverable.",
            task_type="inspection_report",
            selected_models_json='["qwen3.5:4b", "gemma4:e4b"]',
            status="running",
        )
        db.add(task)
        db.commit()

        user_id = user.id
        task_id = task.id

    return {
        "pdf_path": pdf_path,
        "user_id": user_id,
        "task_id": task_id,
    }


def test_demo_a_inspection_workflow_end_to_end(setup_demo_a_environment):
    """Run full Demo A pipeline and verify all sections, citations, calculations, and review gate."""
    from tools.vision import set_vision_hook

    canned_vision = """
    [
      {
        "component": "Circumferential Weld Seam CW-3",
        "visible_condition": "Surface irregularity and localized metal loss visible in heat-affected zone",
        "source": "page_2_image_1.png",
        "limitation": "Visual observation only; not a certified dimensional measurement",
        "type": "observed",
        "confidence": "high"
      }
    ]
    """

    def mock_vision_hook(model, prompt, system, image_b64):
        return canned_vision

    env = setup_demo_a_environment
    pdf_path = env["pdf_path"]
    user_id = env["user_id"]
    task_id = env["task_id"]

    set_vision_hook(mock_vision_hook)
    try:
        result = run_inspection_workflow(
            file_path=pdf_path,
            user_id=user_id,
            role="engineer",
            clearance="CONFIDENTIAL",
            run_id=task_id,
        )
    finally:
        set_vision_hook(None)

    # 1. Output deliverable file existence and checksum
    artifact_p = Path(result.artifact_path)
    assert artifact_p.exists()
    assert artifact_p.stat().st_size > 0
    assert artifact_p.name == "Inspection_Approval_Note.docx"
    assert len(result.sha256) == 64

    # 2. Review gate enforcement (agent CANNOT set APPROVED)
    assert result.approval_status == "PENDING_REVIEW"

    # 3. Output validation checklist
    assert len(result.validation_results) > 0
    failed_validations = [vr for vr in result.validation_results if not vr.passed]
    assert len(failed_validations) == 0, f"Validation failures: {failed_validations}"

    # 4. Verify Calculations performed deterministically
    assert len(result.calculations) >= 3
    calc_params = [c["parameter"] for c in result.calculations]
    assert "General Wall Thinning Loss" in calc_params
    assert "Thinning Threshold Exceedance" in calc_params
    assert "Post-Repair Hydrostatic Proof Pressure" in calc_params

    thinning_calc = next(c for c in result.calculations if "General Wall Thinning" in c["parameter"])
    assert "1.90 mm" in thinning_calc["result"]

    exceed_calc = next(c for c in result.calculations if "Exceedance" in c["parameter"])
    assert "0.40 mm" in exceed_calc["result"]

    proof_calc = next(c for c in result.calculations if "Proof Pressure" in c["parameter"])
    assert "24.75 MPa" in proof_calc["result"]

    # 5. Inspect generated Word deliverable
    doc = docx.Document(artifact_p)
    all_texts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                all_texts.append(cell.text)
    full_doc_text = "\n".join(all_texts)

    # AI-Assisted banner
    assert "AI-ASSISTED DRAFT" in full_doc_text

    # Metadata & Asset
    assert "PV-402" in full_doc_text
    assert "Primary Catalytic Distillation Vessel" in full_doc_text

    # Findings citing sources
    assert "Verified Field Findings (Facts)" in full_doc_text
    assert "demo_scanned_inspection_report.pdf, Page 1" in full_doc_text

    # Visual observations with limitations
    assert "Visual Observations (Multimodal Findings)" in full_doc_text
    assert "Visual observation only; not a certified dimensional measurement" in full_doc_text
    assert "observed" in full_doc_text

    # SOP citations
    assert "SOP-301" in full_doc_text

    # Recommendations separated from facts
    assert "Engineering Recommendations" in full_doc_text
    assert "Mandatory decertification and operational tag-out" in full_doc_text
    assert "24.75 MPa" in full_doc_text

    # Blank Human Review Gate
    assert "Human Review & Sign-Off Gate" in full_doc_text
    assert "Review Decision:" in full_doc_text
    assert "[   ] APPROVED    [   ] REJECTED    [   ] CHANGES REQUESTED" in full_doc_text

    # 6. Check database record
    factory = get_session_factory()
    with factory() as db:
        art_record = db.query(Artifact).filter(Artifact.id == result.artifact_id).first()
        assert art_record is not None
        assert art_record.status == "PENDING_REVIEW"
        assert art_record.reviewer_id is None
