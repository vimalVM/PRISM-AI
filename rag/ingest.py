"""Document ingestion pipeline, parsers, and vector persistence for Sovereign AI Workbench.

Implements 02_DESIGN_DOC.md §8.4 and SEC-04/SEC-05:
- Multi-format parsing: PDF (PyMuPDF), DOCX (python-docx), XLSX (openpyxl), PPTX (python-pptx), TXT/MD.
- Air-gapped embeddings via local SentenceTransformer on CPU (Qwen3-Embedding-0.6B).
- Persistent ChromaDB indexing (data/chroma) with telemetry strictly disabled.
- Automated version superseding (marking old version chunks superseded=True for audit).
- SQLite kb_documents tracking and cryptographic audit logging (kb_ingested, kb_deleted).
"""

from datetime import datetime, timezone
import hashlib
import os
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple, Union

import chromadb
from chromadb.config import Settings as ChromaSettings
from sqlalchemy import select

from backend.core.audit import log_event
from backend.core.config import get_settings
from backend.core.db import KBDocument, get_session_factory
from backend.core.paths import AccessDenied, safe_path
from models.registry import get_registry
from rag.chunking import DocumentChunk, chunk_table, chunk_text


# ChromaDB collection name
KB_COLLECTION_NAME = "kb_chunks"

_CHROMA_CLIENT: Optional[chromadb.PersistentClient] = None
_EMBEDDING_MODEL: Optional[Any] = None
_EMBEDDING_HOOK: Optional[Any] = None  # Hook for mocking in fast tests


def set_embedding_hook(hook: Optional[Any]) -> None:
    """Override embedding generator for fast testing without model weights."""
    global _EMBEDDING_HOOK
    _EMBEDDING_HOOK = hook


def get_chroma_client(chroma_dir: Optional[Union[str, Path]] = None) -> chromadb.PersistentClient:
    """Get persistent ChromaDB client with telemetry strictly disabled."""
    global _CHROMA_CLIENT
    if _CHROMA_CLIENT is None or chroma_dir is not None:
        target_dir = str(chroma_dir) if chroma_dir else "data/chroma"
        os.makedirs(target_dir, exist_ok=True)
        settings = ChromaSettings(
            anonymized_telemetry=False,
            is_persistent=True,
            persist_directory=target_dir,
        )
        client = chromadb.PersistentClient(path=target_dir, settings=settings)
        if chroma_dir is None:
            _CHROMA_CLIENT = client
        return client
    return _CHROMA_CLIENT


def get_kb_collection(chroma_dir: Optional[Union[str, Path]] = None):
    """Retrieve or create the kb_chunks collection."""
    client = get_chroma_client(chroma_dir)
    return client.get_or_create_collection(
        name=KB_COLLECTION_NAME,
        metadata={"description": "Sovereign AI Workbench Knowledge Base Chunks"},
    )


def get_embedding_model() -> Any:
    """Load offline SentenceTransformer model on CPU."""
    global _EMBEDDING_MODEL, _EMBEDDING_HOOK
    if _EMBEDDING_HOOK is not None:
        return _EMBEDDING_HOOK

    if _EMBEDDING_MODEL is None:
        os.environ["HF_HUB_OFFLINE"] = "1"
        from sentence_transformers import SentenceTransformer

        settings = get_settings()
        registry = get_registry(settings.MODEL_REGISTRY_PATH)
        model_path = registry.embeddings.model_path
        device = registry.embeddings.device or "cpu"

        _EMBEDDING_MODEL = SentenceTransformer(model_path, device=device)
    return _EMBEDDING_MODEL


def compute_embeddings(texts: List[str]) -> List[List[float]]:
    """Compute dense vector embeddings for input texts."""
    if not texts:
        return []

    global _EMBEDDING_HOOK
    if _EMBEDDING_HOOK is not None:
        return _EMBEDDING_HOOK(texts)

    model = get_embedding_model()
    embeddings = model.encode(texts, normalize_embeddings=True, show_progress_bar=False)
    return [e.tolist() for e in embeddings]


# --- File Parsers ---

def parse_pdf(file_path: Path) -> List[Tuple[int, str]]:
    """Extract text page-by-page from PDF using PyMuPDF (fitz)."""
    import fitz

    doc = fitz.open(str(file_path))
    pages: List[Tuple[int, str]] = []
    for page_idx in range(len(doc)):
        page = doc[page_idx]
        text = page.get_text("text")
        pages.append((page_idx + 1, text))
    doc.close()
    return pages


def parse_docx(file_path: Path) -> Tuple[List[Tuple[int, str]], List[Tuple[int, List[str], List[List[str]]]]]:
    """Extract paragraphs and tables from DOCX using python-docx."""
    import docx

    doc = docx.Document(str(file_path))
    paragraphs: List[str] = []
    for p in doc.paragraphs:
        if p.text.strip():
            paragraphs.append(p.text.strip())

    text_pages = [(1, "\n\n".join(paragraphs))]

    tables: List[Tuple[int, List[str], List[List[str]]]] = []
    for t_idx, table in enumerate(doc.tables):
        if not table.rows:
            continue
        headers = [cell.text.strip() for cell in table.rows[0].cells]
        rows = []
        for row in table.rows[1:]:
            rows.append([cell.text.strip() for cell in row.cells])
        tables.append((1, headers, rows))

    return text_pages, tables


def parse_xlsx(file_path: Path) -> List[Tuple[int, str, List[str], List[List[str]]]]:
    """Extract sheets and tabular data from XLSX using openpyxl."""
    import openpyxl

    wb = openpyxl.load_workbook(str(file_path), data_only=True)
    sheets_data: List[Tuple[int, str, List[str], List[List[str]]]] = []

    for sheet_idx, sheet_name in enumerate(wb.sheetnames, start=1):
        ws = wb[sheet_name]
        all_rows = list(ws.iter_rows(values_only=True))
        if not all_rows:
            continue

        # Treat first non-empty row as header
        header_row_idx = 0
        while header_row_idx < len(all_rows) and not any(all_rows[header_row_idx]):
            header_row_idx += 1

        if header_row_idx >= len(all_rows):
            continue

        headers = [str(c) if c is not None else "" for c in all_rows[header_row_idx]]
        rows = []
        for r in all_rows[header_row_idx + 1 :]:
            if any(c is not None for c in r):
                rows.append([str(c) if c is not None else "" for c in r])

        sheets_data.append((sheet_idx, sheet_name, headers, rows))

    wb.close()
    return sheets_data


def parse_pptx(file_path: Path) -> List[Tuple[int, str]]:
    """Extract text slide-by-slide from PPTX using python-pptx."""
    from pptx import Presentation

    prs = Presentation(str(file_path))
    slides: List[Tuple[int, str]] = []

    for slide_idx, slide in enumerate(prs.slides, start=1):
        slide_texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    if para.text.strip():
                        slide_texts.append(para.text.strip())
            elif shape.has_table:
                for row in shape.table.rows:
                    row_txt = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_txt:
                        slide_texts.append(row_txt)
        slides.append((slide_idx, "\n".join(slide_texts)))

    return slides


def parse_plain_text(file_path: Path) -> List[Tuple[int, str]]:
    """Read plain text or markdown file."""
    try:
        content = file_path.read_text(encoding="utf-8")
    except UnicodeDecodeError:
        content = file_path.read_text(encoding="latin-1", errors="replace")
    return [(1, content)]


# --- Ingestion Entrypoint ---

def ingest_document(
    file_path: Union[str, Path],
    doc_id: str,
    classification: int = 1,
    version: int = 1,
    uploaded_by: str = "system",
    chroma_dir: Optional[Union[str, Path]] = None,
) -> Dict[str, Any]:
    """Ingest a document into local Knowledge Base with versioning and access control metadata.

    1. Parse source document (PDF, DOCX, XLSX, PPTX, TXT, MD).
    2. Split content into heading-aware chunks and tables.
    3. If previous version of doc_id exists, mark its chunks superseded=True.
    4. Compute CPU embeddings and upsert into ChromaDB kb_chunks.
    5. Update SQLite kb_documents and emit kb_ingested audit event.
    """
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"Source file not found at: {path}")

    content_bytes = path.read_bytes()
    file_sha256 = hashlib.sha256(content_bytes).hexdigest()
    filename = path.name
    ext = path.suffix.lower()

    chunks: List[DocumentChunk] = []

    # 1. Parse by extension
    if ext == ".pdf":
        pages = parse_pdf(path)
        for page_num, text in pages:
            chunks.extend(
                chunk_text(
                    text=text,
                    doc_id=doc_id,
                    filename=filename,
                    doc_version=version,
                    classification=classification,
                    page=page_num,
                    initial_section=f"Page {page_num}",
                )
            )
    elif ext in {".docx", ".doc"}:
        text_pages, tables = parse_docx(path)
        for page_num, text in text_pages:
            chunks.extend(
                chunk_text(
                    text=text,
                    doc_id=doc_id,
                    filename=filename,
                    doc_version=version,
                    classification=classification,
                    page=page_num,
                )
            )
        for page_num, headers, rows in tables:
            chunks.extend(
                chunk_table(
                    headers=headers,
                    rows=rows,
                    doc_id=doc_id,
                    filename=filename,
                    doc_version=version,
                    classification=classification,
                    page=page_num,
                )
            )
    elif ext in {".xlsx", ".xls"}:
        sheets = parse_xlsx(path)
        for sheet_idx, sheet_name, headers, rows in sheets:
            chunks.extend(
                chunk_table(
                    headers=headers,
                    rows=rows,
                    doc_id=doc_id,
                    filename=filename,
                    doc_version=version,
                    classification=classification,
                    page=sheet_idx,
                    section=sheet_name,
                )
            )
    elif ext in {".pptx", ".ppt"}:
        slides = parse_pptx(path)
        for slide_num, text in slides:
            chunks.extend(
                chunk_text(
                    text=text,
                    doc_id=doc_id,
                    filename=filename,
                    doc_version=version,
                    classification=classification,
                    page=slide_num,
                    initial_section=f"Slide {slide_num}",
                )
            )
    else:  # Plain text, Markdown, CSV, Python, etc.
        pages = parse_plain_text(path)
        for page_num, text in pages:
            chunks.extend(
                chunk_text(
                    text=text,
                    doc_id=doc_id,
                    filename=filename,
                    doc_version=version,
                    classification=classification,
                    page=page_num,
                )
            )

    if not chunks:
        # Fallback single chunk for minimal empty documents
        chunks.append(
            DocumentChunk(
                chunk_id=f"{doc_id}_v{version}_p1_c1",
                doc_id=doc_id,
                filename=filename,
                page=1,
                section="General",
                doc_version=version,
                classification=classification,
                text=path.read_text(encoding="utf-8", errors="replace")[:1000],
                content_hash=file_sha256,
                superseded=False,
            )
        )

    # 2. Access Vector DB collection
    collection = get_kb_collection(chroma_dir)

    # 3. Handle versioning: supersede old chunks for the same doc_id
    existing_records = collection.get(where={"doc_id": doc_id})
    if existing_records and existing_records.get("ids"):
        old_ids = existing_records["ids"]
        old_metadatas = existing_records.get("metadatas", [])
        updated_metadatas = []
        for m in old_metadatas:
            updated_m = dict(m)
            updated_m["superseded"] = True
            updated_metadatas.append(updated_m)
        collection.update(ids=old_ids, metadatas=updated_metadatas)

    # 4. Generate embeddings and upsert new chunks
    chunk_texts = [c.text for c in chunks]
    embeddings = compute_embeddings(chunk_texts)

    ids = [c.chunk_id for c in chunks]
    metadatas = [
        {
            "doc_id": c.doc_id,
            "filename": c.filename,
            "page": c.page,
            "section": c.section,
            "doc_version": c.doc_version,
            "classification": c.classification,
            "content_hash": c.content_hash,
            "superseded": c.superseded,
            "is_table": c.is_table,
        }
        for c in chunks
    ]

    collection.upsert(
        ids=ids,
        embeddings=embeddings,
        documents=chunk_texts,
        metadatas=metadatas,
    )

    # 5. Persist to database
    factory = get_session_factory()
    with factory() as db:
        from backend.core.db import User
        user_row = db.execute(select(User).where(User.id == uploaded_by)).scalar_one_or_none()
        actual_uploader = uploaded_by
        if not user_row:
            sys_user = db.execute(select(User).where(User.username == "system")).scalar_one_or_none()
            if not sys_user:
                sys_user = User(
                    id="system",
                    username="system",
                    password_hash="system_locked_account",
                    role="admin",
                    clearance=3,
                    active=True,
                )
                db.add(sys_user)
                db.commit()
            actual_uploader = sys_user.id

        # Mark previous document version superseded
        stmt = select(KBDocument).where(KBDocument.doc_id == doc_id, KBDocument.status == "active")
        prev_docs = db.execute(stmt).scalars().all()
        for pd in prev_docs:
            pd.status = "superseded"

        kb_doc = KBDocument(
            doc_id=doc_id,
            filename=filename,
            version=version,
            classification=classification,
            sha256=file_sha256,
            status="active",
            chunks=len(chunks),
            uploaded_by=actual_uploader,
        )
        db.add(kb_doc)
        db.commit()

    # 6. Audit log
    log_event(
        event_type="kb_ingested",
        status="ok",
        user_id=uploaded_by,
        role="admin",
        details={
            "doc_id": doc_id,
            "filename": filename,
            "version": version,
            "classification": classification,
            "chunk_count": len(chunks),
            "sha256": file_sha256,
        },
    )

    return {
        "doc_id": doc_id,
        "filename": filename,
        "version": version,
        "classification": classification,
        "chunks_indexed": len(chunks),
        "sha256": file_sha256,
    }


def delete_document(
    doc_id: str,
    user_id: str = "system",
    chroma_dir: Optional[Union[str, Path]] = None,
) -> bool:
    """Delete a document and all its chunks from the Knowledge Base."""
    collection = get_kb_collection(chroma_dir)
    collection.delete(where={"doc_id": doc_id})

    factory = get_session_factory()
    with factory() as db:
        stmt = select(KBDocument).where(KBDocument.doc_id == doc_id)
        docs = db.execute(stmt).scalars().all()
        for d in docs:
            d.status = "deleted"
        db.commit()

    log_event(
        event_type="kb_deleted",
        status="ok",
        user_id=user_id,
        role="admin",
        details={"doc_id": doc_id},
    )
    return True
